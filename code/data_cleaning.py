"""
data_cleaning.py
================================================================
Loads the final event-timing workbook, confirms the final sample
(38 Bank of England press conferences), maps cleaned transcript files
to meetings, and loads the underlying one-minute market data.

Two meetings with edited/prerecorded formats (2020 Q4, 2021 Q1) are
excluded upstream of this script and are not present in the input files.

NOTE: The step that removes journalists' questions from the raw Q&A
transcript (retaining only Bank of England policymakers' answers) is
not included here — it was applied prior to the files this script reads.
================================================================
"""
import re
import numpy as np
import pandas as pd
from pathlib import Path
from datetime import datetime, timedelta, time
import openpyxl
import warnings

warnings.simplefilter("ignore")

DATA_DIR = Path("data")
OUTPUT_DIR = Path("outputs")

EVENT_MASTER_FILE = DATA_DIR / "event_master.xlsx"
MARKET_DATA_FILE = DATA_DIR / "market_data_1min.xlsx"
PR_DIR = DATA_DIR / "clean_texts" / "prepared_remarks"
QA_DIR = DATA_DIR / "clean_texts" / "qa_policymaker_only"

ASSETS = ["GBPUSD", "FTSE", "UK2Y", "UK10Y"]

# Bloomberg sheet name -> internal instrument key, and internal key -> asset prefix
INSTRUMENTS = {"GBPUSD": "GBPUSD", "Gilt_2Y": "UK_2Y", "Gilt_10Y": "UK_10Y", "UKX": "FTSE100"}
SEC_PREFIX = {"GBPUSD": "GBPUSD", "FTSE100": "FTSE", "UK_2Y": "UK2Y", "UK_10Y": "UK10Y"}

# GBP/USD and FTSE use log returns; gilt yields use level (percentage-point) changes
LOG_RETURN_ASSETS = {"GBPUSD", "FTSE100"}

POLICY_SURPRISE_VARS = ["MPC_Target", "MPC_Path", "MPC_QE"]


# ================================================================
# Final sample construction
# ================================================================

def load_event_timing():
    """Final event-timing sheets:
        Boundary_Unbuffered - exact Prepared Remarks / Q&A section boundaries,
                              used for the direct PR-vs-Q&A repricing comparison
        Boundary_Buffered   - whole-conference window (PR start - 5 min to
                              Q&A end + 10 min), used for the regression outcome
    """
    timing_unbuffered = pd.read_excel(EVENT_MASTER_FILE, sheet_name="Boundary_Unbuffered", header=1)
    timing_buffered = pd.read_excel(EVENT_MASTER_FILE, sheet_name="Boundary_Buffered", header=1)
    meeting_master = pd.read_excel(EVENT_MASTER_FILE, sheet_name="Meeting_Master", header=1)

    assert len(meeting_master) == 38, f"Expected 38 meetings, found {len(meeting_master)}"
    assert meeting_master["Meeting_ID"].duplicated().sum() == 0, "Duplicate Meeting_ID found"

    for col in ["PR_Start_Rounded_Time", "PR_End_Rounded_Time", "QA_Start_Rounded_Time", "QA_End_Rounded_Time"]:
        timing_unbuffered[col] = pd.to_datetime(timing_unbuffered[col])

    timing_unbuffered["PR_Duration_Min"] = (
        timing_unbuffered["PR_End_Rounded_Time"] - timing_unbuffered["PR_Start_Rounded_Time"]
    ).dt.total_seconds() / 60
    timing_unbuffered["QA_Duration_Min"] = (
        timing_unbuffered["QA_End_Rounded_Time"] - timing_unbuffered["QA_Start_Rounded_Time"]
    ).dt.total_seconds() / 60

    return meeting_master, timing_unbuffered, timing_buffered


def load_policy_surprises():
    """Announcement-window monetary-policy surprise factors (MPC_Target,
    MPC_Path, MPC_QE). Two meetings (2020_Q2, 2020_Q3) are missing these
    and are excluded from the controlled regression sample (N=36)."""
    policy_df = pd.read_excel(EVENT_MASTER_FILE, sheet_name="Meeting_Master", header=1)[
        ["Meeting_ID"] + POLICY_SURPRISE_VARS
    ]
    missing_meetings = policy_df.loc[policy_df[POLICY_SURPRISE_VARS].isna().any(axis=1), "Meeting_ID"].tolist()
    print(f"Meetings missing policy-surprise data: {missing_meetings}")
    return policy_df


# ================================================================
# Transcript file mapping (Prepared Remarks and Q&A supplied separately)
# ================================================================

def get_meeting_id(file_path):
    match = re.search(r"BOE_(\d{4})_Q([1-4])", file_path.name)
    return f"{match.group(1)}_Q{match.group(2)}" if match else None


def build_file_map(folder):
    return {get_meeting_id(p): p for p in folder.glob("*.txt") if get_meeting_id(p) is not None}


def load_transcript_files():
    pr_files = build_file_map(PR_DIR)
    qa_files = build_file_map(QA_DIR)
    meeting_ids = sorted(set(pr_files.keys()) & set(qa_files.keys()))
    assert len(meeting_ids) == 38
    return pr_files, qa_files, meeting_ids


# ================================================================
# One-minute market data loading
# ================================================================

def parse_meeting_id_full(mid):
    match = re.match(r"(?:IR|MPR)-(\d{4})-Q(\d)", mid)
    return f"{match.group(1)}_Q{match.group(2)}" if match else None


def parse_intraday_sheet(ws):
    """Parse a Bloomberg-format sheet into one timestamp-price series per
    meeting block."""
    block_re = re.compile(r"\s*(\d+)\.\s+(\S+)\s+(\d{4}-\d{2}-\d{2})")
    blocks = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str):
            m = block_re.match(v)
            if m:
                blocks.append({"row": r, "mid": m.group(2), "date": m.group(3)})
    out = {}
    for bi, blk in enumerate(blocks):
        start = blk["row"] + 1
        end = blocks[bi + 1]["row"] if bi + 1 < len(blocks) else ws.max_row + 1
        mdate = datetime.strptime(blk["date"], "%Y-%m-%d").date()
        rows = []
        for r in range(start, end):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if isinstance(a, datetime) and a.date() == mdate and b is not None:
                try:
                    rows.append((a, float(b)))
                except (TypeError, ValueError):
                    pass
        if rows:
            out[parse_meeting_id_full(blk["mid"])] = (
                pd.DataFrame(rows, columns=["Timestamp", "Price"]).sort_values("Timestamp").reset_index(drop=True)
            )
    return out


def load_market_data():
    wb_market = openpyxl.load_workbook(MARKET_DATA_FILE, data_only=True)
    return {key: parse_intraday_sheet(wb_market[sheet]) for sheet, key in INSTRUMENTS.items()}


if __name__ == "__main__":
    meeting_master, timing_unbuffered, timing_buffered = load_event_timing()
    pr_files, qa_files, meeting_ids = load_transcript_files()
    intraday = load_market_data()
    print(f"Final sample: N = {len(meeting_ids)} meetings")
